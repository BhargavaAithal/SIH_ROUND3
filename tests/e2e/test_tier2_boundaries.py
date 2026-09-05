"""
Tier 2: Boundary & Corner Invariant Test Suite (25 Discrete Tests)
Validates extreme, non-nominal, and adversarial inputs for R1 through R5
per explorer_e2e_1/analysis.md.
"""
import math
import os
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
import zipfile
import docx
import networkx as nx
import numpy as np
import openpyxl
import pytest

from sovereign.agent import ImmutableSpec, LoopOutcome, run_react_loop
from sovereign.reports import generate_audit_workbook, generate_psu_memo
from sovereign.sandbox import audit_network_egress, run_sandboxed, SandboxResult
from sovereign.verifier import (
    ASTVerificationResult,
    verify_api_510_invariants,
    verify_asme_b31_3,
    verify_python_ast,
)
from sovereign.vision import (
    extract_topology,
    find_junctions_and_endpoints,
    Patch,
    skeletonize_lines,
    slice_drawing,
)


# ========================================================================================
# FEATURE R1 BOUNDARY & CORNER CASES (5 TESTS)
# ========================================================================================

@pytest.mark.tier2
@pytest.mark.boundary_r1
def test_r1_boundary_zero_and_negative_timeout(sandbox_runner):
    """Test 2.B1.1: Error raised when timeout_sec <= 0."""
    with pytest.raises(ValueError):
        sandbox_runner(["echo", "1"], timeout_sec=0)
    with pytest.raises(ValueError):
        sandbox_runner(["echo", "1"], timeout_sec=-5)


@pytest.mark.tier2
@pytest.mark.boundary_r1
def test_r1_boundary_extreme_memory_limit_low(sandbox_runner):
    """Test 2.B1.2: Pathological memory constraint (<10MB) cleanly contained."""
    result = sandbox_runner([sys.executable, "-c", "print(1)"], memory_limit_mb=5)
    assert result.returncode != 0
    assert "MemoryError" in result.stderr or result.returncode != 0


@pytest.mark.tier2
@pytest.mark.boundary_r1
def test_r1_boundary_fork_bomb_process_limit(sandbox_runner):
    """Test 2.B1.3: Recursive fork or thread creation is cleanly terminated."""
    code = "import os; [getattr(os, 'fork', lambda: 1)() for _ in range(20)]"
    result = sandbox_runner([sys.executable, "-c", code], timeout_sec=3)
    # Process terminates cleanly without hanging the test runner
    assert isinstance(result.returncode, int)


@pytest.mark.tier2
@pytest.mark.boundary_r1
def test_r1_boundary_audit_corrupted_or_missing_log():
    """Test 2.B1.4: Missing or invalid audit log handled gracefully."""
    non_existent = "/tmp/non_existent_airgap_log_12345.log"
    res = audit_network_egress(non_existent)
    assert isinstance(res, dict) or hasattr(res, "status")
    assert "status" in res or "verdict" in res
    assert res.get("egress_detected") is False or res["egress_detected"] is False


@pytest.mark.tier2
@pytest.mark.boundary_r1
def test_r1_boundary_sanitized_environment_secrets_leak(monkeypatch, sandbox_runner):
    """Test 2.B1.5: Host secrets are stripped from sandboxed environment."""
    monkeypatch.setenv("SUPER_SECRET_TOKEN", "CLASSIFIED_12345")
    code = "import os; print(os.environ.get('SUPER_SECRET_TOKEN', 'NOT_FOUND'))"
    result = sandbox_runner([sys.executable, "-c", code])
    assert "CLASSIFIED_12345" not in result.stdout
    assert "NOT_FOUND" in result.stdout


# ========================================================================================
# FEATURE R2 BOUNDARY & CORNER CASES (5 TESTS)
# ========================================================================================

@pytest.mark.tier2
@pytest.mark.boundary_r2
def test_r2_boundary_empty_and_uniform_image():
    """Test 2.B2.1: Blank uniform images produce empty graph without division by zero."""
    blank_white = np.ones((3000, 4000), dtype=np.uint8) * 255
    blank_black = np.zeros((3000, 4000), dtype=np.uint8)

    skel_w = skeletonize_lines(blank_white)
    skel_b = skeletonize_lines(blank_black)
    assert np.count_nonzero(skel_w) == 0
    assert np.count_nonzero(skel_b) == 0

    j_w, ep_w = find_junctions_and_endpoints(skel_w)
    assert j_w == [] and ep_w == []

    graph = extract_topology(blank_white)
    assert graph.number_of_nodes() == 0


@pytest.mark.tier2
@pytest.mark.boundary_r2
def test_r2_boundary_disconnected_subgraphs_isolation():
    """Test 2.B2.2: Independent process lines do not cross-contaminate."""
    # Image with 3 disjoint lines
    img = np.full((1500, 2000), 255, dtype=np.uint8)
    import cv2
    cv2.line(img, (200, 300), (800, 300), (0, 0, 0), 4)
    cv2.line(img, (200, 750), (800, 750), (0, 0, 0), 4)
    cv2.line(img, (200, 1200), (800, 1200), (0, 0, 0), 4)

    tags = [
        {"tag": "TAG-1", "type": "pump", "centroid": (200, 300)},
        {"tag": "TAG-2", "type": "valve", "centroid": (200, 750)},
        {"tag": "TAG-3", "type": "exchanger", "centroid": (200, 1200)},
    ]
    graph = extract_topology(img, tags_data=tags)
    assert "TAG-1" in graph.nodes
    assert "TAG-2" in graph.nodes
    assert "TAG-3" in graph.nodes
    assert nx.has_path(graph, "TAG-1", "TAG-2") is False
    assert nx.has_path(graph, "TAG-2", "TAG-3") is False


@pytest.mark.tier2
@pytest.mark.boundary_r2
def test_r2_boundary_dense_junction_noise_and_crossing():
    """Test 2.B2.3: Salt-and-pepper noise does not cause false junction explosion."""
    img = np.full((1000, 1000), 255, dtype=np.uint8)
    import cv2
    # Crossing lines
    cv2.line(img, (200, 500), (800, 500), (0, 0, 0), 5)
    cv2.line(img, (500, 200), (500, 800), (0, 0, 0), 5)

    # 40 small noise speckles
    np.random.seed(42)
    for _ in range(40):
        rx, ry = np.random.randint(50, 950, size=2)
        cv2.circle(img, (int(rx), int(ry)), 1, (0, 0, 0), -1)

    skel = skeletonize_lines(img)
    junctions, endpoints = find_junctions_and_endpoints(skel)
    # The crossing at (500, 500) detected without spurious explosions
    assert len(junctions) <= 4
    assert len(endpoints) <= 6


@pytest.mark.tier2
@pytest.mark.boundary_r2
def test_r2_boundary_extreme_aspect_ratios():
    """Test 2.B2.4: Non-square panoramic (8000x600) and column (600x6000) tiling."""
    pano = np.zeros((600, 8000), dtype=np.uint8)
    column = np.zeros((6000, 600), dtype=np.uint8)

    patches_p = slice_drawing(pano, tile_size=(1024, 1024), overlap=128)
    patches_c = slice_drawing(column, tile_size=(1024, 1024), overlap=128)

    assert len(patches_p) >= 8
    assert len(patches_c) >= 6
    assert max(p.x_offset + p.width for p in patches_p) == 8000
    assert max(p.y_offset + p.height for p in patches_c) == 6000


@pytest.mark.tier2
@pytest.mark.boundary_r2
def test_r2_boundary_equipment_tag_split_across_patch_seam():
    """Test 2.B2.5: Tag centroid on seam boundary is deduplicated."""
    img = np.full((2000, 2000), 255, dtype=np.uint8)
    # Tag centroid at seam X=1024
    tags = [
        {"tag": "10-P-101-CS", "type": "pump", "centroid": (1024, 1000)},
        {"tag": "10-P-101-CS", "type": "pump", "centroid": (1024, 1000)},  # duplicate
    ]
    graph = extract_topology(img, tags_data=tags)
    matching = [n for n in graph.nodes if n == "10-P-101-CS"]
    assert len(matching) == 1


# ========================================================================================
# FEATURE R3 BOUNDARY & CORNER CASES (5 TESTS)
# ========================================================================================

@pytest.mark.tier2
@pytest.mark.boundary_r3
def test_r3_boundary_zero_and_negative_pressure():
    """Test 2.B3.1: Zero and negative pressure rejected under ASME B31.3 Section 304.1.2."""
    res_zero = verify_asme_b31_3(
        design_pressure=0.0,
        outside_diameter=323.8,
        allowable_stress=137.9,
        quality_factor=1.0,
        temp_coefficient=0.4,
        corrosion_allowance=3.0,
        actual_thickness=9.52,
    )
    res_neg = verify_asme_b31_3(
        design_pressure=-2.5,
        outside_diameter=323.8,
        allowable_stress=137.9,
        quality_factor=1.0,
        temp_coefficient=0.4,
        corrosion_allowance=3.0,
        actual_thickness=9.52,
    )
    assert res_zero.is_valid is False and res_zero.status == "UNSAT"
    assert res_neg.is_valid is False and res_neg.status == "UNSAT"
    assert any("P must be > 0" in v or "pressure" in v.lower() for v in res_zero.violations)


@pytest.mark.tier2
@pytest.mark.boundary_r3
def test_r3_boundary_zero_and_negative_thickness():
    """Test 2.B3.2: Non-physical zero or negative thickness rejected."""
    res_zero = verify_asme_b31_3(
        design_pressure=2.5,
        outside_diameter=323.8,
        allowable_stress=137.9,
        quality_factor=1.0,
        temp_coefficient=0.4,
        corrosion_allowance=3.0,
        actual_thickness=0.0,
    )
    res_neg = verify_api_510_invariants(t_actual=-5.0, t_min=10.0, pressure=3.5)
    assert res_zero.is_valid is False and res_zero.status == "UNSAT"
    assert res_neg.is_valid is False and res_neg.status == "UNSAT"


@pytest.mark.tier2
@pytest.mark.boundary_r3
def test_r3_boundary_extreme_temperature_coefficients():
    """Test 2.B3.3: Invariant bounds for Y in [0, 0.7] and E in [0.6, 1.0]."""
    res1 = verify_asme_b31_3(
        design_pressure=2.5,
        outside_diameter=323.8,
        allowable_stress=137.9,
        quality_factor=1.2,  # out of bounds
        temp_coefficient=0.4,
        corrosion_allowance=3.0,
        actual_thickness=9.52,
    )
    res2 = verify_asme_b31_3(
        design_pressure=2.5,
        outside_diameter=323.8,
        allowable_stress=137.9,
        quality_factor=1.0,
        temp_coefficient=1.5,  # out of bounds
        corrosion_allowance=3.0,
        actual_thickness=9.52,
    )
    assert res1.is_valid is False
    assert res2.is_valid is False


@pytest.mark.tier2
@pytest.mark.boundary_r3
def test_r3_boundary_ast_obfuscation_and_introspection():
    """Test 2.B3.4: Adversarial reflection / obfuscation patterns rejected."""
    code1 = "getattr(__builtins__, 'ex' + 'ec')('import os; os.system(\"echo 1\")')"
    code2 = "globals()['__builtins__']['eval']('__import__(\"socket\")')"
    res1 = verify_python_ast(code1)
    res2 = verify_python_ast(code2)
    assert res1.is_safe is False
    assert res2.is_safe is False


@pytest.mark.tier2
@pytest.mark.boundary_r3
def test_r3_boundary_ast_syntax_error_handling():
    """Test 2.B3.5: Malformed Python code does not crash verifier."""
    code = "def broken_calc(x: return x +"
    res = verify_python_ast(code)
    assert res.is_safe is False
    assert any("SyntaxError" in v for v in res.violations)


# ========================================================================================
# FEATURE R4 BOUNDARY & CORNER CASES (5 TESTS)
# ========================================================================================

@pytest.mark.tier2
@pytest.mark.boundary_r4
def test_r4_boundary_immediate_success_turn_one():
    """Test 2.B4.1: Script passing on Turn 1 exits immediately without extra turns."""
    spec = ImmutableSpec(
        task_id="TASK-T1-PASS",
        task_description="Immediate success test"
    )

    def valid_engine(prompt: str) -> str:
        return "import json; print(json.dumps([{'verdict': 'SAT'}]))"

    outcome = run_react_loop(spec, max_turns=3, engine_callback=valid_engine)
    assert outcome.success is True
    assert outcome.turns_used == 1
    assert len(outcome.failure_hashes) == 0


@pytest.mark.tier2
@pytest.mark.boundary_r4
def test_r4_boundary_sandbox_timeout_handling_in_loop():
    """Test 2.B4.2: Infinite loop caught, categorized as TIMEOUT, and self-corrected in Turn 2."""
    spec = ImmutableSpec(
        task_id="TASK-TIMEOUT-RECOVER",
        task_description="Recover from infinite loop"
    )

    calls = [0]

    def timeout_recovery_engine(prompt: str) -> str:
        calls[0] += 1
        if calls[0] == 1:
            return "import time; time.sleep(10)"  # times out
        return "import json; print(json.dumps([{'status': 'SAT'}]))"

    outcome = run_react_loop(spec, max_turns=3, engine_callback=timeout_recovery_engine)
    assert outcome.turns_used == 2
    assert outcome.success is True
    assert outcome.trace[0]["category"] == "TIMEOUT"


@pytest.mark.tier2
@pytest.mark.boundary_r4
def test_r4_boundary_ast_blocked_script_reprompting():
    """Test 2.B4.3: AST rejection prevents execution and reprompts model."""
    spec = ImmutableSpec(
        task_id="TASK-AST-RECOVER",
        task_description="Recover from forbidden import"
    )

    calls = [0]

    def ast_recovery_engine(prompt: str) -> str:
        calls[0] += 1
        if calls[0] == 1:
            return "import requests; print('leak')"
        return "import json; print(json.dumps([{'status': 'SAT'}]))"

    outcome = run_react_loop(spec, max_turns=3, engine_callback=ast_recovery_engine)
    assert outcome.turns_used == 2
    assert outcome.success is True
    assert "requests" in outcome.trace[0]["error"]


@pytest.mark.tier2
@pytest.mark.boundary_r4
def test_r4_boundary_empty_or_none_spec_validation():
    """Test 2.B4.4: Reject null or unpopulated ImmutableSpec."""
    with pytest.raises(ValueError):
        run_react_loop(None)
    with pytest.raises(ValueError):
        run_react_loop(ImmutableSpec(task_id="", task_description=""))


@pytest.mark.tier2
@pytest.mark.boundary_r4
def test_r4_boundary_identical_repeated_failure_shortcircuit():
    """Test 2.B4.5: Repeated identical failure shortcircuits to prevent token waste."""
    spec = ImmutableSpec(
        task_id="TASK-STALL",
        task_description="Detect identical failure stalls"
    )

    def stuck_engine(prompt: str) -> str:
        return "raise RuntimeError('Identical Stalled Error')"

    outcome = run_react_loop(spec, max_turns=3, engine_callback=stuck_engine)
    assert outcome.success is False
    assert any(entry.get("stalled") is True for entry in outcome.trace)


# ========================================================================================
# FEATURE R5 BOUNDARY & CORNER CASES (5 TESTS)
# ========================================================================================

@pytest.mark.tier2
@pytest.mark.boundary_r5
def test_r5_boundary_large_dataset_table_pagination(tmp_path):
    """Test 2.B5.1: 10,000 calculation rows generated without memory exhaustion."""
    xlsx_path = str(tmp_path / "Large_Audit.xlsx")
    large_dataset = [
        {"Point": i, "P": 2.5, "D": 323.8, "t_act": 9.52, "verdict": "SAT"}
        for i in range(1, 10002)
    ]
    generate_audit_workbook({"LargeRun": large_dataset}, xlsx_path)
    assert os.path.exists(xlsx_path) and os.path.getsize(xlsx_path) > 50000
    wb = openpyxl.load_workbook(xlsx_path)
    assert wb["LargeRun"].max_row >= 10001


@pytest.mark.tier2
@pytest.mark.boundary_r5
def test_r5_boundary_special_characters_and_xml_injection(tmp_path):
    """Test 2.B5.2: XML reserved chars and unicode symbols render without corruption."""
    docx_path = str(tmp_path / "Special_Chars.docx")
    xlsx_path = str(tmp_path / "Special_Chars.xlsx")

    meta = {
        "title": "TEST <SAFETY> & 'RISK' \"ASSESSMENT\" -- 100% AIRGAP",
        "ref_no": "PSU/REF/<001>&'26'",
        "engineer": "श्रीमान शर्मा & Dr. O'Connor",
        "tag": "10-P-101 <HIGH-PRESS>",
    }
    calc = [
        {"tag": "VALVE <HV-101> & 'SAFE'", "notes": "Tested at 120°C & π*R² with Ω resistance", "verdict": "SAT"}
    ]

    generate_psu_memo(meta, calc, ["Code: ASME B31.3 <Edition 2022> & OISD-141"], docx_path)
    generate_audit_workbook({"Special": calc}, xlsx_path)

    with zipfile.ZipFile(docx_path, 'r') as z:
        assert z.testzip() is None
        doc_xml = z.read("word/document.xml")
        assert ET.fromstring(doc_xml) is not None

    with zipfile.ZipFile(xlsx_path, 'r') as z:
        assert z.testzip() is None
        wb_xml = z.read("xl/workbook.xml")
        assert ET.fromstring(wb_xml) is not None


@pytest.mark.tier2
@pytest.mark.boundary_r5
def test_r5_boundary_empty_data_structures_handling(tmp_path):
    """Test 2.B5.3: Empty metadata and empty calculation lists handle cleanly."""
    docx_path = str(tmp_path / "Empty_Memo.docx")
    xlsx_path = str(tmp_path / "Empty_Workbook.xlsx")

    generate_psu_memo({}, [], [], docx_path)
    generate_audit_workbook({}, xlsx_path)

    assert os.path.exists(docx_path)
    assert os.path.exists(xlsx_path)

    with zipfile.ZipFile(docx_path, 'r') as z:
        assert z.testzip() is None
    with zipfile.ZipFile(xlsx_path, 'r') as z:
        assert z.testzip() is None


@pytest.mark.tier2
@pytest.mark.boundary_r5
def test_r5_boundary_file_permission_and_path_errors():
    """Test 2.B5.4: Unwritable or invalid path raises clean error."""
    invalid_path = "/nonexistent_root_dir_abc123/unwritable/memo.docx"
    with pytest.raises((OSError, FileNotFoundError, PermissionError)):
        generate_psu_memo({}, [], [], invalid_path)


@pytest.mark.tier2
@pytest.mark.boundary_r5
def test_r5_boundary_extreme_numerical_precision_and_overflow(tmp_path):
    """Test 2.B5.5: NaN and extreme floats safely serialized without workbook crash."""
    xlsx_path = str(tmp_path / "Extreme_Numbers.xlsx")
    data = [
        {"Tag": "EXT-1", "Value": 1.5e9, "Verdict": "SAT"},
        {"Tag": "EXT-2", "Value": 1.23456789e-7, "Verdict": "SAT"},
        {"Tag": "EXT-3", "Value": float("nan"), "Verdict": "UNKNOWN"},
        {"Tag": "EXT-4", "Value": float("inf"), "Verdict": "INVALID"},
    ]
    generate_audit_workbook({"Extreme": data}, xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["Extreme"]
    val_nan = ws["B4"].value
    assert val_nan in ("N/A", "#NUM!", None)
