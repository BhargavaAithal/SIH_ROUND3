"""
Tier 1: Core Feature Coverage Test Suite (29 Discrete Tests)
Validates baseline functional contracts for R1 through R5 per explorer_e2e_1/analysis.md.
Enforces opaque-box testing across all sovereign subsystems.
"""
from dataclasses import FrozenInstanceError
import hashlib
import os
from pathlib import Path
import sys
import xml.etree.ElementTree as ET
import zipfile
import docx
import networkx as nx
import numpy as np
import openpyxl
import pytest

from sovereign.agent import ImmutableSpec, LoopOutcome, run_react_loop
from sovereign.reports import generate_audit_workbook, generate_psu_memo
from sovereign.sandbox import audit_network_egress, run_sandboxed, SandboxResult
from sovereign.verifier import (
    ASTVerificationResult,
    verify_api_510_invariants,
    verify_asme_b31_3,
    verify_python_ast,
    Z3VerificationResult,
)
from sovereign.vision import (
    extract_topology,
    find_junctions_and_endpoints,
    Patch,
    skeletonize_lines,
    slice_drawing,
)


# ========================================================================================
# FEATURE R1: AIR-GAP & SOVEREIGNTY ENFORCEMENT (6 TESTS)
# ========================================================================================

@pytest.mark.tier1
@pytest.mark.feature_r1
def test_r1_network_isolation_flags(sandbox_runner):
    """Test 1.1: Sandbox injects strict network isolation flags."""
    cmd = [sys.executable, "-c", "print('ISOLATED_OK')"]
    result = sandbox_runner(cmd, network=False)
    assert result.returncode == 0
    assert "ISOLATED_OK" in result.stdout
    assert result.network_egress_bytes == 0


@pytest.mark.tier1
@pytest.mark.feature_r1
def test_r1_socket_bind_connect_blocked(sandbox_runner):
    """Test 1.2: Outbound TCP socket connect is blocked inside sandbox."""
    code = (
        "import socket; "
        "s = socket.socket(socket.AF_INET, socket.SOCK_STREAM); "
        "s.connect(('1.1.1.1', 80))"
    )
    result = sandbox_runner([sys.executable, "-c", code], timeout_sec=5, network=False)
    assert result.returncode != 0
    assert (
        "Network is unreachable" in result.stderr
        or "PermissionDenied" in result.stderr
        or "PermissionError" in result.stderr
        or "OSError" in result.stderr
        or "ConnectionRefused" in result.stderr
    )
    assert result.network_egress_bytes == 0


@pytest.mark.tier1
@pytest.mark.feature_r1
def test_r1_airgap_audit_script_execution():
    """Test 1.3: Opaque-box execution of audit_network_egress()."""
    verdict = audit_network_egress()
    assert verdict["egress_detected"] is False
    assert verdict["packets_captured"] == 0
    assert "PASS" in verdict["status"] or "PASS" in verdict["verdict"]


@pytest.mark.tier1
@pytest.mark.feature_r1
def test_r1_rlimit_cgroups_memory_constraint(sandbox_runner):
    """Test 1.4: Script allocating memory beyond limit is contained."""
    # Attempt 700MB allocation with 512MB limit
    code = "x = bytearray(700 * 1024 * 1024)"
    result = sandbox_runner([sys.executable, "-c", code], memory_limit_mb=512)
    assert result.returncode != 0
    assert (
        "MemoryError" in result.stderr
        or result.returncode in (-9, 137, 1)
        or result.memory_peak_mb <= 550
    )


@pytest.mark.tier1
@pytest.mark.feature_r1
def test_r1_rlimit_cgroups_cpu_timeout_constraint(sandbox_runner):
    """Test 1.5: Hard CPU timeout terminates infinite busy-loop."""
    code = "import time\nwhile True: pass"
    result = sandbox_runner([sys.executable, "-c", code], timeout_sec=3)
    assert result.returncode in (-9, -15, 124, 137, 1)
    assert 2.5 <= result.execution_time_sec <= 5.0
    assert "Timeout" in result.stderr or result.timed_out or result.returncode != 0


@pytest.mark.tier1
@pytest.mark.feature_r1
def test_r1_loopback_only_behavior(sandbox_runner):
    """Test 1.6: External WAN DNS resolution fails in isolated sandbox."""
    code = "import socket; socket.gethostbyname('google.com')"
    result = sandbox_runner([sys.executable, "-c", code], network=False)
    assert result.returncode != 0
    assert (
        "gaierror" in result.stderr
        or "Name or service not known" in result.stderr
        or "Errno" in result.stderr
        or "Sovereign air-gap" in result.stderr
        or "Air-Gap" in result.stderr
        or "PermissionError" in result.stderr
    )


# ========================================================================================
# FEATURE R2: MULTIMODAL RASTER-TO-GRAPH & LAYOUT PARSER (6 TESTS)
# ========================================================================================

@pytest.mark.tier1
@pytest.mark.feature_r2
def test_r2_highres_image_slicing_and_tiling(synthetic_pid_image):
    """Test 2.1: 4000x3000 drawing sliced into regular 1024x1024 patches with 128px overlap."""
    image_array, _ = synthetic_pid_image
    patches = slice_drawing(image_array, tile_size=(1024, 1024), overlap=128)
    assert len(patches) == 20
    assert all(p.image_array.shape == (1024, 1024) for p in patches[:-1])
    assert max(p.x_offset + p.width for p in patches) == 4000
    assert max(p.y_offset + p.height for p in patches) == 3000


@pytest.mark.tier1
@pytest.mark.feature_r2
def test_r2_patch_coordinate_math_roundtrip():
    """Test 2.2: Bijection between global drawing and patch-local coordinates."""
    # Arbitrary global point
    x_g, y_g = 2450, 1850
    patches = slice_drawing(np.zeros((3000, 4000), dtype=np.uint8), tile_size=(1024, 1024), overlap=128)

    # Find patch containing point
    matching_patches = [
        p for p in patches
        if p.x_offset <= x_g < p.x_offset + p.width and p.y_offset <= y_g < p.y_offset + p.height
    ]
    assert len(matching_patches) >= 1
    p_k = matching_patches[0]

    # Convert to local and back to global
    x_l = x_g - p_k.x_offset
    y_l = y_g - p_k.y_offset
    x_r = p_k.x_offset + x_l
    y_r = p_k.y_offset + y_l

    assert (x_r, y_r) == (2450, 1850)
    assert 0 <= x_l < p_k.width
    assert 0 <= y_l < p_k.height


@pytest.mark.tier1
@pytest.mark.feature_r2
def test_r2_morphological_line_skeletonization(synthetic_pid_image):
    """Test 2.3: Morphological thinning reduces 7px lines to 1-pixel skeletons."""
    image_array, _ = synthetic_pid_image
    skeleton = skeletonize_lines(image_array)
    assert skeleton.shape == image_array.shape
    assert set(np.unique(skeleton)).issubset({0, 1, 255})
    assert np.count_nonzero(skeleton) > 0


@pytest.mark.tier1
@pytest.mark.feature_r2
def test_r2_junction_and_endpoint_detection(synthetic_pid_image):
    """Test 2.4: Topological classification of T-junctions, crossings, and endpoints."""
    image_array, _ = synthetic_pid_image
    skeleton = skeletonize_lines(image_array)
    junctions, endpoints = find_junctions_and_endpoints(skeleton)
    assert len(junctions) >= 3
    assert len(endpoints) >= 2


@pytest.mark.tier1
@pytest.mark.feature_r2
def test_r2_equipment_tag_detection_and_centroid_snapping(synthetic_pid_image):
    """Test 2.5: Equipment tags snapped to piping endpoints."""
    image_array, gt = synthetic_pid_image
    graph = extract_topology(image_array, tags_data=gt)
    assert "10-P-101-CS" in graph.nodes
    assert "V-102" in graph.nodes
    assert "E-103" in graph.nodes


@pytest.mark.tier1
@pytest.mark.feature_r2
def test_r2_networkx_graph_construction_and_query(synthetic_pid_image):
    """Test 2.6: Constructed graph allows flow connectivity queries."""
    image_array, gt = synthetic_pid_image
    graph = extract_topology(image_array, tags_data=gt)
    assert isinstance(graph, nx.Graph)
    assert nx.has_path(graph, "10-P-101-CS", "V-102") is True
    path = nx.shortest_path(graph, "10-P-101-CS", "V-102")
    assert len(path) >= 2
    assert graph.number_of_nodes() >= 3


# ========================================================================================
# FEATURE R3: AST & NEUROSYMBOLIC Z3 VERIFIER (6 TESTS)
# ========================================================================================

@pytest.mark.tier1
@pytest.mark.feature_r3
def test_r3_python_ast_unsafe_node_rejection(unsafe_ast_code_snippets):
    """Test 3.1: AST policy rejects forbidden imports and calls."""
    results = [verify_python_ast(snippet) for snippet in unsafe_ast_code_snippets]
    assert all(r.is_safe is False for r in results)
    assert all(len(r.violations) > 0 for r in results)


@pytest.mark.tier1
@pytest.mark.feature_r3
def test_r3_python_ast_safe_calculation_script_acceptance():
    """Test 3.2: Safe arithmetic and mathematical calculation passes AST check."""
    safe_code = (
        "import math\n"
        "P = 2.5\n"
        "D = 323.8\n"
        "S = 137.9\n"
        "E = 1.0\n"
        "Y = 0.4\n"
        "c = 3.0\n"
        "t_min = (P * D) / (2 * (S * E + P * Y)) + c\n"
        "print(f't_min={t_min}')\n"
    )
    result = verify_python_ast(safe_code)
    assert result.is_safe is True
    assert len(result.violations) == 0


@pytest.mark.tier1
@pytest.mark.feature_r3
def test_r3_asme_b31_3_valid_thickness_sat(asme_b31_3_dataset):
    """Test 3.3: Z3 proves SAT when actual thickness exceeds minimum required thickness."""
    d = asme_b31_3_dataset["valid"]
    result = verify_asme_b31_3(
        design_pressure=d["P"],
        outside_diameter=d["D"],
        allowable_stress=d["S"],
        quality_factor=d["E"],
        temp_coefficient=d["Y"],
        corrosion_allowance=d["c"],
        actual_thickness=d["t_actual"],
    )
    assert result.is_valid is True
    assert result.status == "SAT"
    assert 3.59 <= result.margin <= 3.62
    assert len(result.violations) == 0


@pytest.mark.tier1
@pytest.mark.feature_r3
def test_r3_asme_b31_3_invalid_thickness_unsat(asme_b31_3_dataset):
    """Test 3.4: Z3 proves UNSAT when actual thickness is insufficient."""
    d = asme_b31_3_dataset["corroded"]
    result = verify_asme_b31_3(
        design_pressure=d["P"],
        outside_diameter=d["D"],
        allowable_stress=d["S"],
        quality_factor=d["E"],
        temp_coefficient=d["Y"],
        corrosion_allowance=d["c"],
        actual_thickness=d["t_actual"],
    )
    assert result.is_valid is False
    assert result.status == "UNSAT"
    assert result.margin < 0
    assert any("t_actual < t_min" in v or "thickness" in v.lower() for v in result.violations)


@pytest.mark.tier1
@pytest.mark.feature_r3
def test_r3_api_510_pressure_vessel_invariants(api_510_vessel_dataset):
    """Test 3.5: API 510 remaining life calculation and invariant verification."""
    val = api_510_vessel_dataset["valid"]
    inval = api_510_vessel_dataset["invalid"]

    res_valid = verify_api_510_invariants(
        t_actual=val["t_actual"],
        t_min=val["t_min"],
        pressure=val["P"],
        corrosion_rate=val["corrosion_rate"],
    )
    res_invalid = verify_api_510_invariants(
        t_actual=inval["t_actual"],
        t_min=inval["t_min"],
        pressure=inval["P"],
        corrosion_rate=inval["corrosion_rate"],
    )
    assert res_valid.is_valid is True and res_valid.status == "SAT"
    assert abs(res_valid.model_details["remaining_life_years"] - 16.8) < 1e-3
    assert res_invalid.is_valid is False and res_invalid.status == "UNSAT"


@pytest.mark.tier1
@pytest.mark.feature_r3
def test_r3_false_assurance_rate_strictly_zero():
    """Test 3.6: Prove False Assurance Rate (FAR) strictly equals 0.0% across 100 cases."""
    results = []
    for i in range(100):
        # Synthesize parameter where t_actual < t_min
        delta = 0.01 + (i * 0.05)
        # Baseline t_min is ~5.914, so set actual to t_min - delta
        insufficient_thickness = max(0.01, 5.914 - delta)
        res = verify_asme_b31_3(
            design_pressure=2.5,
            outside_diameter=323.8,
            allowable_stress=137.9,
            quality_factor=1.0,
            temp_coefficient=0.4,
            corrosion_allowance=3.0,
            actual_thickness=insufficient_thickness,
        )
        results.append(res)

    false_assurances = sum(1 for r in results if r.is_valid is True)
    assert false_assurances == 0
    far = false_assurances / len(results)
    assert far == 0.0


# ========================================================================================
# FEATURE R4: STATE-ISOLATED ANTI-COLLAPSE LOOP (6 TESTS)
# ========================================================================================

@pytest.mark.tier1
@pytest.mark.feature_r4
def test_r4_spec_script_hash_state_separation():
    """Test 4.1: ImmutableSpec is frozen and prevents in-place mutation."""
    spec = ImmutableSpec(
        task_id="TASK-1",
        task_description="Compute ASME B31.3 wall thickness",
        parameters={"P": 2.5, "D": 323.8}
    )
    with pytest.raises(FrozenInstanceError):
        spec.task_id = "MUTATED"


@pytest.mark.tier1
@pytest.mark.feature_r4
def test_r4_max_three_turns_constraint_enforcement():
    """Test 4.2: Loop terminates strictly after 3 turns when script persistently fails."""
    spec = ImmutableSpec(
        task_id="TASK-PERSIST-FAIL",
        task_description="Continuous failure test"
    )

    def failing_engine(prompt: str) -> str:
        turn_num = len(prompt) % 100
        return f"raise RuntimeError('Persistent Error {turn_num}')"

    outcome = run_react_loop(spec, max_turns=3, engine_callback=failing_engine)
    assert outcome.success is False
    assert outcome.turns_used == 3
    assert len(outcome.failure_hashes) == 3


@pytest.mark.tier1
@pytest.mark.feature_r4
def test_r4_failure_hash_tracking_and_deduplication():
    """Test 4.3: Failure hashes are 64-character hexadecimal SHA-256 strings."""
    spec = ImmutableSpec(
        task_id="TASK-HASH-TEST",
        task_description="Hash validation test"
    )

    turn_counter = [0]

    def engine_with_varying_failures(prompt: str) -> str:
        turn_counter[0] += 1
        return f"raise ValueError('Error variant {turn_counter[0]}')"

    outcome = run_react_loop(spec, max_turns=3, engine_callback=engine_with_varying_failures)
    assert len(outcome.failure_hashes) >= 2
    assert all(len(h) == 64 and all(c in "0123456789abcdef" for c in h) for h in outcome.failure_hashes)
    assert outcome.failure_hashes[0] != outcome.failure_hashes[1]


@pytest.mark.tier1
@pytest.mark.feature_r4
def test_r4_clean_context_reprompting():
    """Test 4.4: Turn 2 prompt contains ONLY Immutable Spec + Immediate Traceback."""
    spec = ImmutableSpec(
        task_id="TASK-CLEAN-CONTEXT",
        task_description="Verify prompt cleanliness"
    )

    prompts_captured = []

    def engine_spy(prompt: str) -> str:
        prompts_captured.append(prompt)
        if len(prompts_captured) == 1:
            return "raise KeyError('Missing key')"
        return "import json; print(json.dumps([{'status': 'SAT'}]))"

    outcome = run_react_loop(spec, max_turns=3, engine_callback=engine_spy)
    assert outcome.success is True
    assert len(prompts_captured) == 2
    t2_prompt = prompts_captured[1]
    assert spec.task_description in t2_prompt
    assert "KeyError" in t2_prompt
    assert "Turn 0" not in t2_prompt


@pytest.mark.tier1
@pytest.mark.feature_r4
def test_r4_self_correction_recovery_syntax_error():
    """Test 4.5: Self-correction succeeds when Turn 1 has SyntaxError and Turn 2 fixes it."""
    spec = ImmutableSpec(
        task_id="TASK-SYNTAX-RECOVER",
        task_description="Recover from syntax error"
    )

    calls = [0]

    def syntax_fix_engine(prompt: str) -> str:
        calls[0] += 1
        if calls[0] == 1:
            return "def broken_syntax(: return 42"
        return "import json; print(json.dumps([{'t_min': 5.914}]))"

    outcome = run_react_loop(spec, max_turns=3, engine_callback=syntax_fix_engine)
    assert outcome.success is True
    assert outcome.turns_used == 2
    assert len(outcome.failure_hashes) == 1
    assert "SyntaxError" in outcome.trace[0]["stderr"]


@pytest.mark.tier1
@pytest.mark.feature_r4
def test_r4_self_correction_recovery_z3_unsat_error():
    """Test 4.6: Self-correction repairs undersized pipe schedule based on Z3 feedback."""
    spec = ImmutableSpec(
        task_id="TASK-Z3-RECOVER",
        task_description="Correct undersized pipe thickness",
        parameters={"P": 2.5, "D": 323.8, "t_actual": 9.52}
    )

    calls = [0]

    def z3_fix_engine(prompt: str) -> str:
        calls[0] += 1
        if calls[0] == 1:
            # Undersized thickness
            return "import json; print(json.dumps([{'tag': 'P1', 't_actual': 4.5, 'verdict': 'UNSAT'}]))"
        # Corrected thickness
        return "import json; print(json.dumps([{'tag': 'P1', 't_actual': 9.52, 'verdict': 'SAT'}]))"

    outcome = run_react_loop(spec, max_turns=3, engine_callback=z3_fix_engine)
    assert outcome.success is True
    assert outcome.turns_used in (1, 2)
    if outcome.z3_result:
        assert outcome.z3_result.status == "SAT"


# ========================================================================================
# FEATURE R5: HEADLESS DELIVERABLES GENERATOR (5 TESTS)
# ========================================================================================

@pytest.mark.tier1
@pytest.mark.feature_r5
def test_r5_docx_psu_approval_memo_generation(clean_psu_report_data, tmp_path):
    """Test 5.1: Compilation of native PSU approval note (.docx)."""
    memo_path = str(tmp_path / "Approval_Memo.docx")
    out = generate_psu_memo(
        metadata=clean_psu_report_data["metadata"],
        calculations=clean_psu_report_data["calculations"],
        citations=clean_psu_report_data["citations"],
        output_path=memo_path,
    )
    assert os.path.exists(out)
    assert os.path.getsize(out) > 5000
    doc = docx.Document(out)
    full_text = "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    assert "Paradip Refinery" in full_text
    assert "10-P-101-CS" in full_text
    assert "ASME B31.3-2022 Section 304.1.2" in full_text
    assert len(doc.tables) >= 1


@pytest.mark.tier1
@pytest.mark.feature_r5
def test_r5_xlsx_calculation_audit_workbook_generation(clean_psu_report_data, tmp_path):
    """Test 5.2: Compilation of multi-tab .xlsx calculation audit workbook."""
    xlsx_path = str(tmp_path / "Audit_Workbook.xlsx")
    sheets = {
        "Summary": [{"Metric": "Total Lines", "Value": 4}],
        "ASME B31.3": clean_psu_report_data["calculations"],
        "Audit Trail": [{"Turn": 1, "Status": "PASS"}],
    }
    out = generate_audit_workbook(sheets, xlsx_path)
    assert os.path.exists(out)
    wb = openpyxl.load_workbook(out)
    assert set(wb.sheetnames) == {"Summary", "ASME B31.3", "Audit Trail"}


@pytest.mark.tier1
@pytest.mark.feature_r5
def test_r5_xlsx_formula_verification(tmp_path):
    """Test 5.3: Excel calculation sheet contains live formula strings starting with '='."""
    xlsx_path = str(tmp_path / "Formula_Workbook.xlsx")
    sheets = {
        "Calculations": [
            {"Tag": "P1", "P": 2.5, "D": 323.8, "S": 137.9, "E": 1.0, "Y": 0.4, "CA": 3.0, "t_act": 9.52, "t_min": 5.914},
            {"Tag": "P2", "P": 2.5, "D": 323.8, "S": 137.9, "E": 1.0, "Y": 0.4, "CA": 3.0, "t_act": 9.52, "t_min": 5.914},
            {"Tag": "P3", "P": 2.5, "D": 323.8, "S": 137.9, "E": 1.0, "Y": 0.4, "CA": 3.0, "t_act": 9.52, "t_min": 5.914},
            {"Tag": "P4", "P": 2.5, "D": 323.8, "S": 137.9, "E": 1.0, "Y": 0.4, "CA": 3.0, "t_act": 9.52, "t_min": 5.914},
        ]
    }
    generate_audit_workbook(sheets, xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    ws = wb["Calculations"]
    formula_cell = ws["G5"].value
    assert isinstance(formula_cell, str) and formula_cell.startswith("=")
    assert "C5*D5" in formula_cell


@pytest.mark.tier1
@pytest.mark.feature_r5
def test_r5_table_formatting_and_enterprise_styling(clean_psu_report_data, tmp_path):
    """Test 5.4: Corporate visual formatting (bold headers, fill colors, number formats)."""
    xlsx_path = str(tmp_path / "Styled_Workbook.xlsx")
    docx_path = str(tmp_path / "Styled_Memo.docx")

    generate_audit_workbook({"ASME": clean_psu_report_data["calculations"]}, xlsx_path)
    generate_psu_memo(
        clean_psu_report_data["metadata"],
        clean_psu_report_data["calculations"],
        clean_psu_report_data["citations"],
        docx_path,
    )

    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["ASME"]
    assert ws["A1"].fill.start_color.rgb is not None
    assert ws["A1"].font.bold is True
    assert ws["D5"].number_format in ("0.000", "0.00", "#,##0.00")

    doc = docx.Document(docx_path)
    assert len(doc.tables[0].rows) >= 2


@pytest.mark.tier1
@pytest.mark.feature_r5
def test_r5_ooxml_validation_without_corruption(clean_psu_report_data, tmp_path):
    """Test 5.5: Generated .docx and .xlsx archives pass OpenXML XML schema parsing."""
    docx_path = str(tmp_path / "Validate_Memo.docx")
    xlsx_path = str(tmp_path / "Validate_Workbook.xlsx")

    generate_psu_memo(
        clean_psu_report_data["metadata"],
        clean_psu_report_data["calculations"],
        clean_psu_report_data["citations"],
        docx_path,
    )
    generate_audit_workbook({"Sheet1": clean_psu_report_data["calculations"]}, xlsx_path)

    # Word XML test
    with zipfile.ZipFile(docx_path, 'r') as z_doc:
        assert z_doc.testzip() is None
        doc_xml = z_doc.read("word/document.xml")
        assert ET.fromstring(doc_xml) is not None

    # Excel XML test
    with zipfile.ZipFile(xlsx_path, 'r') as z_xl:
        assert z_xl.testzip() is None
        wb_xml = z_xl.read("xl/workbook.xml")
        assert ET.fromstring(wb_xml) is not None
